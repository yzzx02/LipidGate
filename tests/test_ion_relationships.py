import json

import numpy as np
import pandas as pd
import pytest

from lipidgate.gui.eic_trace import EICTrace
from lipidgate.ion_relationships import (
    ION_MASSES, compare_peak_shapes, export_ion_relationships,
    infer_ion_relationships, relationship_lines,
)


def features(mode='positive', samples=('a.mzML',)):
    ions = list(ION_MASSES[mode].items())[:3]
    rows=[]
    for i,(ion,mass) in enumerate(ions):
        for sample in samples:
            rows.append(dict(Feature_ID=f'n{i}-{sample}',Aligned_Feature_ID=f'F{i}',
                source_file=sample,mz=500+mass,RT=2.0,RTmin=1.96,RTmax=2.04,
                intensity=(i+1)*100,charge=1))
    return pd.DataFrame(rows)


class Reader:
    def __init__(self, shifted=False, failing=False):
        self.shifted=shifted
        self.failing=failing
        self.calls=[]

    def read_trace(self,path,mz,rt,ppm,window):
        if self.failing:raise ValueError('fixture read failure')
        self.calls.append((path,mz,rt,ppm,window))
        times=np.linspace(rt-window,rt+window,31)
        center=rt+.03 if self.shifted and mz>520 else rt
        signal=np.exp(-.5*((times-center)/.02)**2)*(1+mz)
        return EICTrace(times,signal,rt,mz,ppm,window)


@pytest.mark.parametrize('mode',['positive','negative'])
def test_relationship_hypotheses_preserve_inputs_and_consistent_neutral_mass(mode):
    native=features(mode,samples=('a.mzML','b.mzML'))
    original=native.copy(deep=True)
    reader=Reader()
    relations=infer_ion_relationships(native,mode=mode,trace_reader=reader,sources=['a.mzML','b.mzML'])
    assert len(relations)==3
    assert relations.status.eq('EIC_supported').all()
    assert relations.neutral_mass.tolist()==pytest.approx([500]*3)
    assert relations.supported_samples.eq(2).all()
    pd.testing.assert_frame_equal(native,original)
    assert all(call[4]<=.3 for call in reader.calls)
    assert all('native_feature_a' in e for e in json.loads(relations.iloc[0].evidence_json))


def test_no_default_ion_type_and_no_cross_sample_only_pair():
    native=features().iloc[:1]
    assert infer_ion_relationships(native).empty
    native=features()
    native.loc[native.Aligned_Feature_ID.eq('F2'),'source_file']='other.mzML'
    relation=infer_ion_relationships(native)
    assert len(relation)==1 and relation.iloc[0].status=='mass_RT_only'
    assert not relation.feature_b.eq('ALIGN|F2').any()


def test_native_only_absolute_source_keys_match_browser_keys(tmp_path):
    from lipidgate.gui.result_data import normalize_candidates
    native=features().drop(columns='Aligned_Feature_ID')
    native['source_file']=str(tmp_path/'a.mzML')
    relations=infer_ion_relationships(native,trace_reader=Reader(),sources=[tmp_path/'a.mzML'])
    keys=set(normalize_candidates(native)._key)
    assert set(relations.feature_a) | set(relations.feature_b) == keys
    assert relations.status.eq('EIC_supported').all()


def test_neighboring_local_peaks_cannot_become_supported_ions():
    native=features()
    native.loc[native.Aligned_Feature_ID.eq('F2'),['RT','RTmin','RTmax']]=[2.03,2.02,2.05]
    native.loc[native.Aligned_Feature_ID.ne('F2'),'RTmax']=2.015
    relations=infer_ion_relationships(native,trace_reader=Reader(),sources=['a.mzML'])
    sodium=relations.loc[relations.feature_b.eq('ALIGN|F2')]
    assert len(sodium)==2 and sodium.status.eq('rejected').all()
    assert all(json.loads(row.evidence_json)[0]['status']=='local_peak_disagreement' for row in sodium.itertuples())


def test_shape_disagreement_and_shared_zeroes_do_not_create_support():
    relations=infer_ion_relationships(features(),trace_reader=Reader(shifted=True),sources=['a.mzML'])
    assert relations.loc[relations.feature_b.eq('ALIGN|F2'),'status'].eq('rejected').all()
    times=np.linspace(0,1,101)
    x=np.zeros(101);y=x.copy();x[45:48]=[1,2,1];y[50:53]=[1,2,1]
    a=EICTrace(times,x,.5,501,5);b=EICTrace(times,y,.5,523,5)
    assert compare_peak_shapes(a,b)['status']=='insufficient_points'


def test_conflicting_samples_and_missing_raw_data_remain_explicit():
    class MixedReader(Reader):
        def read_trace(self,path,*args):
            self.shifted=str(path).startswith('b')
            return super().read_trace(path,*args)
    native=features(samples=('a.mzML','b.mzML'))
    relations=infer_ion_relationships(native,trace_reader=MixedReader(),sources=['a.mzML','b.mzML'])
    assert relations.loc[relations.feature_b.eq('ALIGN|F2'),'status'].eq('mixed_evidence').all()
    failed=infer_ion_relationships(native,trace_reader=Reader(failing=True),sources=['a.mzML','b.mzML'])
    assert failed.status.eq('mass_RT_only').all()
    assert 'fixture read failure' in failed.iloc[0].evidence_json


def test_units_known_multicharge_and_isotopes_are_respected():
    native=features();native[['RT','RTmin','RTmax']]*=60;native['RT_unit']='seconds'
    assert len(infer_ion_relationships(native))==3
    native.loc[native.Aligned_Feature_ID.eq('F2'),'charge']=2
    assert len(infer_ion_relationships(native))==1
    aligned=pd.DataFrame([dict(Feature_ID=f'F{i}',mz=mz,RT=2.,**{'a.mzML':100})
                          for i,mz in enumerate(features().mz)])
    assert len(infer_ion_relationships(native,aligned))==1
    native['IsotopeWeightNumber']=[0,1,0]
    assert infer_ion_relationships(native).empty


def test_aligned_only_table_is_never_called_eic_supported():
    aligned=pd.DataFrame([dict(Feature_ID='F1',mz=501.007276466621,RT=2.,**{'a.mzML':100}),
                          dict(Feature_ID='F2',mz=522.989220702,RT=2.,**{'a.mzML':50})])
    result=infer_ion_relationships(None,aligned,trace_reader=Reader(),sources=['a.mzML'])
    assert len(result)==1 and result.iloc[0].status=='mass_RT_only'
    assert relationship_lines(result,'ALIGN|F1','[M+Na]+')[0].endswith('与当前库加合物不同')


def test_missing_raw_file_is_recorded_without_invalidating_export(tmp_path):
    result=export_ion_relationships(features(),None,tmp_path/'results',
                                    sources=[tmp_path/'a.mzML'],mode='positive')
    assert result.status.eq('mass_RT_only').all()
    assert 'eic_read_failure' in result.iloc[0].evidence_json
    metadata=json.loads((tmp_path/'results/ion_relationships.json').read_text(encoding='utf-8'))
    assert metadata['sources'][0]['error']


def test_sidecar_export_and_browser_loading_never_change_search_rows(tmp_path):
    from lipidgate.final_results import export_final_results
    from lipidgate.gui.result_data import load_result_bundle
    candidate=pd.DataFrame([dict(Feature_ID='F1',Aligned_Feature_ID='F1',source_file='a.mzML',
        scan_id='scan_1',precursor_mz=501.007276466621,rt_minutes=2.,matched_name='PE(34:1)',
        compound_class='PE',adduct='[M+Na]+',final_score=80.,result_rank=2,
        passed_required_gates=True,ms1_support_status='MS1-supported')])
    original=candidate.copy(deep=True)
    final=export_final_results(candidate,tmp_path/'results',use_ecn=False,export_xlsx=False,export_plots=False)
    csv_before=final.csv_path.read_bytes()
    audit_before=(tmp_path/'results/audit/evaluated.csv').read_bytes()
    result=export_ion_relationships(features(),None,tmp_path/'results',sources=[],mode='positive')
    assert len(result)==3
    bundle=load_result_bundle(final.csv_path)
    assert len(bundle.ion_relationships)==3
    assert bundle.candidates.iloc[0].adduct=='[M+Na]+'
    assert bundle.candidates.iloc[0].final_score==80 and bundle.candidates.iloc[0].result_rank==2
    assert final.csv_path.read_bytes()==csv_before
    assert (tmp_path/'results/audit/evaluated.csv').read_bytes()==audit_before
    pd.testing.assert_frame_equal(candidate,original)
    from lipidgate.gui.result_export_data import export_browser_results
    from openpyxl import load_workbook
    paths=export_browser_results(bundle,tmp_path/'export',csv=True,xlsx=True,plots=False)
    assert tmp_path/'export/ion_relationships.csv' in paths
    book=load_workbook(tmp_path/'export/LipidGate_results.xlsx')
    assert '加合物关系（推断）' in book.sheetnames
    book.close()
