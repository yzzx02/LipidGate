"""Shared result number formats and workbook writer."""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
import pandas as pd

def format_workbook_sheet(worksheet) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    worksheet.freeze_panes = 'E2'
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for row in worksheet:
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='FFFFFF')
            cell.font = Font(name='Arial', size=11, color='000000', bold=cell.row == 1)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value,str) and cell.value.startswith('='):
                cell.data_type = 's'
    for i,cell in enumerate(worksheet[1],start=1):
        width = 100 if cell.value == 'matched_fragments' else 40 if cell.value in {'matched_name','source_file'} else 20
        worksheet.column_dimensions[get_column_letter(i)].width = width
    header_to_index = {cell.value: index for index, cell in enumerate(worksheet[1], start=1)}
    for column_name, number_format in [
        ("rt_minutes", "0.000"),
        ("selected_ms2_rt", "0.000"),
        ("feature_rt", "0.000"),
        ("feature_rtmin", "0.000"),
        ("feature_rtmax", "0.000"),
        ("precursor_mz", "0.0000"),
        ("feature_mz", "0.0000"),
        ("mz", "0.0000"),
        ("ppm_error", "0.00"),
        ("mz_error_to_feature_ppm", "0.00"),
        ("rt_delta_sec", "0.0"),
        ("final_score", "0.00"),
    ]:
        column_index = header_to_index.get(column_name)
        if column_index is None:
            continue
        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=column_index,
            max_col=column_index,
        ):
            row[0].number_format = number_format



def write_workbook(path: Path, sheets: dict[str, pd.DataFrame], *, permission_fallback: bool = False) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True,exist_ok=True)
    def write(target):
        with pd.ExcelWriter(target, engine="openpyxl") as writer:
            for name, frame in sheets.items():
                frame.to_excel(writer, sheet_name=name, index=False)
                format_workbook_sheet(writer.sheets[name])
        return target
    try:
        return write(path)
    except PermissionError:
        if not permission_fallback:
            raise
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        return write(path.with_name(f"{path.stem}_{stamp}{path.suffix}"))
