import json
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest

from lipidgate.project import Project
from lipidgate.final_results import filter_results, export_final_results


def candidates():
    return pd.DataFrame(
        [
            dict(
                matched_name=f"PC({c}:2)",
                compound_class="PC",
                rt_minutes=c / 4 + 0.02,
                ms1_feature_rt_raw_min=c / 4,
                final_score=90,
                result_rank=1,
                precursor_mz=700 + c,
                adduct="[M+H]+",
                source_file="sample.mzML",
                scan_id=f"scan_{c}",
                ms1_support_status="MS1-supported",
                passed_required_gates=True,
                resolution_level="molecular_species",
            )
            for c in (30, 32, 34, 36, 38)
        ]
    )


def test_project_roundtrip_and_inputs_are_not_copied(tmp_path):
    source = tmp_path / "input.mzML"
    source.write_text("input")
    project = Project.open(tmp_path / "project")
    project.add_files([source, source])
    project.settings = {"ms2": {"mode": "positive"}}
    project.save()
    loaded = Project.open(project.root)
    assert loaded.settings == project.settings
    assert loaded.mzml_files() == [source]
    assert not (project.root / source.name).exists()
    other = tmp_path / "notes.txt"
    other.write_text("notes")
    loaded.add_files([other])
    with pytest.raises(ValueError, match="mzML"):
        loaded.mzml_files()


def test_duplicate_sample_names_are_rejected(tmp_path):
    p = Project.open(tmp_path / "project")
    for folder in ("a", "b"):
        f = tmp_path / folder / "sample.mzML"
        f.parent.mkdir()
        f.touch()
        p.add_files([f])
    with pytest.raises(ValueError, match="同名"):
        p.mzml_files()


def test_multiple_ms1_inputs_stage_only_selected_files(tmp_path):
    from lipidgate.ms1.detection import _mzml_input_dir

    files = []
    for name in ("a.mzML", "b.mzML", "unused.mzML"):
        p = tmp_path / name
        p.write_text(name)
        files.append(p)
    with _mzml_input_dir(files[:2]) as staged:
        assert {p.name for p in staged.iterdir()} == {"a.mzML", "b.mzML"}
        assert (staged / "a.mzML").read_text() == "a.mzML"
    assert not staged.exists()
    assert all(p.exists() for p in files)


@pytest.mark.parametrize("use_ecn", [True, False])
@pytest.mark.parametrize("use_score", [True, False])
def test_empty_export_has_usable_workbook(tmp_path, use_ecn, use_score):
    from openpyxl import load_workbook

    result = export_final_results(
        pd.DataFrame(), tmp_path, use_ecn=use_ecn, use_score=use_score
    )
    wb = load_workbook(result.xlsx_path)
    assert wb.sheetnames == ["最终鉴定", "逐谱证据", "RT未判定"]
    assert wb["最终鉴定"].max_column == 17
    assert len(result.data) == 0
    wb.close()


def test_filters_keep_observed_rt_and_do_not_promote_low_confidence(tmp_path):
    data = candidates()
    low = data.iloc[[2]].copy()
    low["ms1_support_status"] = "MS2-only"
    low["ms1_feature_rt_raw_min"] = 99
    low["rt_minutes"] = 8.9
    low["result_rank"] = 2
    data = pd.concat([data, low], ignore_index=True)
    result = export_final_results(data, tmp_path, plot_dpi=100)
    pd.testing.assert_series_equal(result.data.rt_minutes, data.rt_minutes)
    assert result.data.iloc[-1].rt_for_filter_min == 8.9
    assert result.data.iloc[-1].RT_filter_action == "low_confidence_rescued"
    assert result.data.iloc[-1]["置信度"] == "低"
    assert not result.data.iloc[-1].ECN_training_anchor
    assert len(result.plot_paths) == 1 and result.plot_paths[0].exists()
    scatter = pd.read_csv(tmp_path / "plots/PC_ECN_scatter_data.csv")
    assert scatter.display_rt_min.tolist() == [7.5, 8, 8.5, 9, 9.5]


@pytest.mark.parametrize("use_ecn", [True, False])
def test_score_filter_keeps_correct_original_rows(use_ecn):
    data = candidates()
    data.loc[[1, 3], "final_score"] = 10
    evaluated, _, _ = filter_results(data, use_ecn=use_ecn)
    rejected = evaluated.loc[evaluated.RT_filter_action.eq("score_reject")]
    assert set(rejected.scan_id) == {"scan_32", "scan_36"}
    assert len(evaluated) == len(data)
    unfiltered, _, _ = filter_results(data, use_ecn=False, use_score=False)
    assert unfiltered.RT_consistency_pass.all()
    data.final_score = 1
    evaluated, _, _ = filter_results(data, use_ecn=use_ecn)
    assert evaluated.RT_filter_action.eq("score_reject").all()


def test_ecn_can_retain_subclass_with_too_few_fit_points():
    lone = candidates().iloc[[0]].copy()
    lone["matched_name"] = "LPC(22:6)"
    lone["compound_class"] = "LPC"
    lone["scan_id"] = "scan_lpc"
    data = pd.concat([candidates(), lone], ignore_index=True)
    kept, _, _ = filter_results(data, retain_unmodeled=True)
    omitted, _, _ = filter_results(data, retain_unmodeled=False)
    selected = kept.loc[kept.scan_id.eq("scan_lpc")].iloc[0]
    rejected = omitted.loc[omitted.scan_id.eq("scan_lpc")].iloc[0]
    assert selected.RT_filter_action == "retain_unmodeled"
    assert bool(selected.RT_consistency_pass)
    assert rejected.RT_filter_action == "unmodeled"
    assert not bool(rejected.RT_consistency_pass)


def test_results_page_export_generates_ecn_plot_on_demand(tmp_path):
    from lipidgate.gui.result_data import ResultBundle
    from lipidgate.gui.result_export_data import export_browser_results

    result = export_final_results(candidates(), tmp_path / "run/results",
                                  export_xlsx=False, export_plots=False)
    assert result.plot_paths == []
    bundle = ResultBundle.from_frames(result.data, path=result.csv_path)
    outputs = export_browser_results(bundle, tmp_path / "exports", csv=False,
                                     xlsx=False, plots=True, dpi=300)
    assert any(path.name == "PC_ECN.png" and path.is_file() for path in outputs)


def test_plain_workbook_preserves_numeric_values_and_literal_strings(tmp_path):
    from openpyxl import load_workbook
    from lipidgate.ms2.workbook_export import write_workbook

    path = write_workbook(
        tmp_path / "nested/result.xlsx",
        {"结果": pd.DataFrame({"matched_name": ["=1+1"], "final_score": [80.25]})},
    )
    wb = load_workbook(path)
    ws = wb.active
    assert ws["A2"].data_type == "s" and ws["A2"].value == "=1+1"
    assert ws["B2"].data_type == "n" and ws["B2"].value == 80.25
    assert ws.freeze_panes == "E2"
    assert all(c.fill.fgColor.rgb.endswith("FFFFFF") for row in ws for c in row)
    wb.close()


def test_project_pipeline_snapshots_parameters_and_uses_raw_audit(
    tmp_path, monkeypatch
):
    import lipidgate.pipeline as pipeline

    source = tmp_path / "sample.mzML"
    source.write_text('<mzML><cvParam accession="MS:1000129"/><spectrum><cvParam accession="MS:1000511" value="1"/></spectrum></mzML>')
    library = tmp_path / "library.msp"
    library.write_text("fixture")
    project = Project.open(tmp_path / "project")
    project.add_files([source])
    settings = {
        "ms1": {
            "enabled": True,
            "algo": "xcms",
            "params": {"ppm": 10, "peakwidth": [5, 60]},
        },
        "ms2": {
            "mode": "negative",
            "library_path": str(library),
            "top_n": 3,
            "precursor_tolerance_ppm": 10,
            "fragment_tolerance_ppm": 10,
        },
        "filter": {"use_ecn": False, "use_score": False},
    }
    calls = []

    def detect(**kwargs):
        calls.append(kwargs)
        f = tmp_path / "features.csv"
        f.write_text("mz,RT\n734,8.5\n")
        native = tmp_path / "native.csv"
        native.write_text("source_file,mz,RT\nsample.mzML,734,8.5\n")
        return SimpleNamespace(table_path=f, native_table_path=native)

    def search(**kwargs):
        calls.append(kwargs)
        out = kwargs["output_dir"] / "audit"
        out.mkdir(parents=True)
        candidates().to_csv(out / "ms2_candidates.csv", index=False)
        return SimpleNamespace(data=pd.DataFrame())

    monkeypatch.setattr(pipeline, "run_feature_detection_result", detect)
    monkeypatch.setattr(pipeline, "run_ms2_feature_annotation_result", search)
    _, _, final = pipeline.run_project(project.root, settings)
    assert calls[0]["config"]["parameters"]["xcms"]["polarity"] == "negative"
    assert calls[0]["config"]["parameters"]["asari"]["mode"] == "neg"
    assert calls[0]["input_path"] == [source]
    assert calls[1]["fragment_tolerance_ppm"] == 10
    assert calls[1]["min_total_score"] == 0 and calls[1]["export_csv"]
    assert calls[1]["map_to_features"]
    assert calls[1]["feature_table"] == tmp_path / "native.csv"
    assert len(final.data) == 5
    settings["ms2"]["top_n"] = 1
    assert (
        json.loads((final.output_dir.parent / "run_settings.json").read_text())["ms2"][
            "top_n"
        ]
        == 3
    )
    assert (final.output_dir.parent / "provenance.json").exists()
    ms1_metadata = json.loads((final.output_dir.parent / "ms1/feature_table.json").read_text())
    assert ms1_metadata["algorithm"] == "xcms"
    assert Path(ms1_metadata["table"]) == tmp_path / "features.csv"
    assert Path(ms1_metadata["native_table"]) == tmp_path / "native.csv"
    _, _, second = pipeline.run_project(project.root, settings)
    assert second.output_dir != final.output_dir


def test_ms2_only_input_skips_ms1_extraction_but_keeps_search(tmp_path, monkeypatch):
    import lipidgate.pipeline as pipeline

    source = tmp_path / "sample.mzML"
    source.write_text(
        '<mzML><cvParam accession="MS:1000129"/>'
        '<cvParam accession="MS:1000511" value="2"/></mzML>'
    )
    library = tmp_path / "library.msp"
    library.write_text("fixture")
    project = Project.open(tmp_path / "project")
    project.add_files([source])
    settings = {
        "ms1": {"enabled": True, "algo": "pyopenms", "params": {"min_fwhm": 5, "max_fwhm": 60}},
        "ms2": {"mode": "negative", "library_path": str(library),
                "top_n": 3, "precursor_tolerance_ppm": 10, "fragment_tolerance_ppm": 10},
        "filter": {"use_ecn": False, "use_score": False},
    }
    monkeypatch.setattr(pipeline, "run_feature_detection_result",
                        lambda **kwargs: (_ for _ in ()).throw(AssertionError("MS1 extraction called")))
    searched = []

    def search(**kwargs):
        searched.append(kwargs)
        audit = kwargs["output_dir"] / "audit"
        audit.mkdir(parents=True)
        candidates().to_csv(audit / "ms2_candidates.csv", index=False)
        return SimpleNamespace(data=pd.DataFrame())

    monkeypatch.setattr(pipeline, "run_ms2_feature_annotation_result", search)
    feature, _, final = pipeline.run_project(project.root, settings)
    assert feature is None
    assert searched[0]["feature_table"] is None and not searched[0]["map_to_features"]
    assert final.passed_data.ms1_support_status.eq("MS2-only").all()
    audit = pd.read_csv(final.output_dir / "audit" / "evaluated.csv")
    assert audit.ms1_support_reason.eq("no_ms1_scans").all()


def test_wrong_polarity_rejected_before_work(tmp_path):
    from lipidgate.pipeline import check_polarity

    f = tmp_path / "mixed.mzML"
    f.write_text(
        '<mzML><spectrum><cvParam accession="MS:1000130"/>'
        '<cvParam accession="MS:1000129"/></spectrum></mzML>'
    )
    for mode in ("positive", "negative"):
        with pytest.raises(ValueError, match="模式"):
            check_polarity(f, mode)


def test_opposite_polarity_candidates_rejected_even_without_charge():
    from lipidgate.ms2.search import LipidMS2Searcher
    from lipidgate.ms2.models import LibraryRecord, ExperimentalSpectrum

    r = LibraryRecord(1, "PC", "PC(34:1)", "PC(16:0_18:1)", 760, "[M+H]+")
    spectrum = ExperimentalSpectrum("scan", 760, 1, "-", [])
    assert not LipidMS2Searcher._candidate_charge_is_compatible(spectrum, r)
    spectrum.polarity = "+"
    assert LipidMS2Searcher._candidate_charge_is_compatible(spectrum, r)


def test_peak_truth_is_not_a_product_entry_point():
    import importlib.util
    from lipidgate.cli import build_parser

    assert importlib.util.find_spec("lipidgate.peak_truth") is None
    help_text = build_parser().format_help()
    assert "peak-truth" not in help_text
    assert "project-run" in help_text and "filter-results" in help_text


def test_real_project_detection_search_and_export(tmp_path):
    import numpy as np

    oms = pytest.importorskip("pyopenms")
    from lipidgate.ms2.positive_pe_cer import PositivePECer
    from lipidgate.pipeline import run_project

    m = PositivePECer(18, 1, 16, 0)
    exp = oms.MSExperiment()
    for rt in range(100):
        sp = oms.MSSpectrum()
        sp.setRT(float(rt))
        sp.setMSLevel(1)
        sp.setNativeID(f"scan={rt + 1}")
        sp.setType(1)
        instrument = sp.getInstrumentSettings()
        instrument.setPolarity(1)
        sp.setInstrumentSettings(instrument)
        y = 1e6 * np.exp(-0.5 * ((rt - 40) / 5) ** 2)
        sp.set_peaks(
            (
                np.array(
                    [
                        300.0,
                        m.precursor_mz,
                        m.precursor_mz + 1.00335,
                        m.precursor_mz + 2.0067,
                    ]
                ),
                np.array([6000.0, y, 0.4 * y, 0.08 * y]),
            )
        )
        exp.addSpectrum(sp)
        if rt == 40:
            sp = oms.MSSpectrum()
            sp.setRT(40.1)
            sp.setMSLevel(2)
            sp.setNativeID("scan=101")
            sp.setType(1)
            sp.setInstrumentSettings(instrument)
            precursor = oms.Precursor()
            precursor.setMZ(m.precursor_mz)
            precursor.setCharge(1)
            sp.setPrecursors([precursor])
            peaks = sorted((f.mz, 1000.0) for f in m.fragments())
            sp.set_peaks(
                (np.array([p[0] for p in peaks]), np.array([p[1] for p in peaks]))
            )
            exp.addSpectrum(sp)
    source = tmp_path / "sample.mzML"
    oms.MzMLFile().store(str(source), exp)
    library = tmp_path / "tiny.msp"
    lines = [
        f"Name: {m.name}",
        f"PrecursorMZ: {m.precursor_mz}",
        "PrecursorType: [M+H]+",
        "CompoundClass: PE-Cer",
        f"Num Peaks: {len(m.fragments())}",
    ]
    lines += [f'{f.mz} 100 "{f.name}" "{f.fragment_type}"' for f in m.fragments()]
    library.write_text("\n".join(lines) + "\n\n", encoding="utf-8")
    project = Project.open(tmp_path / "project")
    project.add_files([source])
    settings = {
        "ms1": {
            "enabled": True,
            "algo": "pyopenms",
            "params": dict(
                mz_tol=5.0, min_fwhm=5.0, max_fwhm=60.0, noise=1000.0, sn=5.0,
                min_peak_height=900000.0,
            ),
        },
        "ms2": {
            "mode": "positive",
            "library_path": str(library),
            "precursor_tolerance_ppm": 10,
            "fragment_tolerance_ppm": 10,
        },
        "filter": {"use_ecn": False, "use_score": True},
    }
    feature, _, result = run_project(project.root, settings)
    assert feature.row_count > 0
    assert feature.native_table_path.exists()
    assert result.passed_data.matched_name.tolist() == [m.name]
    assert result.passed_data["置信度"].tolist() == ["高"]
    assert result.passed_data.ms1_support_status.tolist() == ["MS1-supported"]
    assert result.passed_data.feature_rt.iloc[0] == pytest.approx(40 / 60, abs=0.05)
    assert result.xlsx_path is None
    assert result.csv_path.exists()
